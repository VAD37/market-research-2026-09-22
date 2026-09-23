"""Generate docs/findings/director-brief-2026-09-23.pptx.

Reads nothing but its own inline content, which mirrors the tables and key
lines of docs/findings/director-brief-2026-09-23.md. Every number here is also
in that file. The .pptx is a generated file: never hand-edit it; fix this
script and re-run.

Run from anywhere:  python docs/method/gen-director-deck.py
Needs: python-pptx.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "findings" / "director-brief-2026-09-23.pptx"

DARK = RGBColor(0x1F, 0x2A, 0x44)
ACCENT = RGBColor(0x0E, 0x7C, 0x86)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BODY = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
ROW_ALT = RGBColor(0xF2, 0xF5, 0xF7)

FOOTER = "Research only — evidence, not a verdict. Pulls dated 2026-09-22."
SOURCE = "Source: docs/findings/director-brief-2026-09-23.md"

W, H = Inches(13.333), Inches(7.5)

# ---------------------------------------------------------------- content

SLIDES = [
    {
        "kind": "title",
        "title": "Demand for brand visibility inside AI assistants",
        "subtitle": "Director brief · file date 2026-09-23 · long form of the executive brief",
        "lines": [
            "Passes 0–9 and 11 done; Pass 10 held by owner decision 2026-09-22 22:40",
            "All 23 hypotheses touched · all six lanes · compiled files only",
        ],
    },
    {
        "kind": "bullets",
        "title": "Agenda",
        "bullets": [
            "The question and the short answer",
            "What was done",
            "The market: size, paid inventory, agentic checkout, regulation",
            "Demand: the 27-cell segment read",
            "Proof: what the published cases show",
            "What brands that moved changed",
            "Hypotheses and programme-done rows",
            "What the evidence does not show",
            "Questions only the owner can answer · caveats",
        ],
    },
    {
        "kind": "bullets",
        "title": "The question",
        "quote": "Which segments show real-world demand for visibility and recommendation inside "
                 "AI assistants, what evidence shows it works, and what did brands that moved "
                 "actually change? (method/plan.md L11)",
        "bullets": [
            "Demand: 7 of 27 cells read spend; 5 enterprise; 6 of 7 rest on job postings",
            "Proof: 0 Gold of ~980 screened; no holdout, geo-split or switchback",
            "Paid: live on ChatGPT and Google, refused by Claude; no rate card",
            "Size: no sub-market measured; every dollar size is a forecast",
            "Change: content ops 63 and stack 58 cases against paid media 2",
            "Own sampling: partial day 0 only; held 2026-09-22 22:40",
        ],
    },
    {
        "kind": "table",
        "title": "What was done",
        "cols": ["Pass", "What it pulled", "Count"],
        "widths": [1.6, 5.4, 5.3],
        "rows": [
            ["0–1", "method rules; source channels, query book", "5 deliverables; 42 clusters after AMEND-1"],
            ["2–3", "platform primaries; vendor census", "~150 raw pulls; 121 screened, 26 rostered, ~190 pulls"],
            ["4", "published success stories", "13 clusters; ~980 candidates screened"],
            ["5", "manipulation techniques, countermeasures", "7 clusters, ~70 raw pulls"],
            ["6–7", "market sizes; competitor profiles", "16 size pulls, 18 figures; 41 profiles + INDEX"],
            ["8", "demand signals per segment", "15 + 10 + 12 signal pulls; 27 cells"],
            ["9, 11", "findings and review; transition evidence", "review recounted 27 claims; 108 cases, 99 name a change"],
            ["10 — held", "our own engine panel, day 0", "Claude 83 of 160; Gemini 24 of 160; AI Mode 76 of 176; ChatGPT, Perplexity 0"],
        ],
    },
    {
        "kind": "table",
        "title": "Market size: none measured, forecasts side by side",
        "key": "No sub-market has a measured size. Forecasts are forecasts, never a size.",
        "cols": ["Metric", "Figure", "Label, tier"],
        "widths": [3.2, 6.3, 2.8],
        "rows": [
            ["Measured size", "none; 15 of 18 published figures are forecasts", "analyst-derived, 6"],
            ["Organic, bottom-up floor", "$5.3M–$35.2M annualised; 4 of 34 vendors (~12%)", "vendor-reported inputs, 5"],
            ["Paid, agentic bottom-up", "not computable from disclosed inputs", "—"],
            ["Forecast, organic", "Valuates US$7,318M 2031 · Market Decipher USD 32.92B 2034", "forecast, 6"],
            ["Forecast, paid", "EMARKETER US $68.25B 2030 · WPP Media global over $100 billion 2030", "forecast"],
            ["Forecast, agentic", "EMARKETER US $144B 2029 · McKinsey global $3T–$5T 2030", "forecast, 6 · 5"],
            ["Agentic spread", "USD 144B to USD 5T, roughly 35×, scopes not comparable", "forecast"],
            ["Attention proxy", "IAB: 76% of >200 buyers name AI answers No. 1 focus", "analyst-derived, 4"],
        ],
    },
    {
        "kind": "table",
        "title": "Paid inventory, per engine",
        "key": "ChatGPT ads: $1 billion annualized run rate (company-stated, tier 3, 2026-08-31). 0 of 8 rate cards.",
        "cols": ["Engine / metric", "Status", "Label, tier"],
        "widths": [3.2, 6.6, 2.5],
        "rows": [
            ["ChatGPT", "live from 2026-02-09; self-serve Ads Manager beta", "company-stated, 3"],
            ["Google AI Overviews", "live; existing campaigns auto-eligible", "company-stated, 3"],
            ["Google AI Mode · Gemini app", "testing · unknown — checked 2026-09-22", "company-stated, 3"],
            ["Claude", "no — \"Claude will remain ad-free\", 2026-02-04", "company-stated, 3"],
            ["Coverage", "5 of 8 engines live; 0 of 8 rate cards; 1 of 8 revenue", "company-stated, 3"],
            ["Only price stated", "OpenAI max-bid guidance $3–$5 USD per click", "company-stated, 3"],
            ["ChatGPT ad presence", "26% · 4.47% US · 0.00% of 169,560 UK scrapes", "vendor-reported, 5"],
            ["Measured by us", "0 ad units in 76 AI Mode runs; 0 of 14 AIO rendered", "measured-by-us, 1"],
        ],
    },
    {
        "kind": "table",
        "title": "Agentic checkout and regulation",
        "cols": ["Item", "Figure", "Label, tier"],
        "widths": [3.2, 6.8, 2.3],
        "rows": [
            ["Live checkout programs", "4 — ChatGPT, Google, Copilot, Perplexity", "company-stated, 3"],
            ["Fee disclosed", "1 of 4 — Copilot, no commission; protocol fee clauses 6 of 6 unknown", "company-stated, 3"],
            ["Open protocols", "ACP, UCP, x402 Apache-2.0 (H8 confirmed)", "company-stated, 3"],
            ["AI conversion", "Adobe 42% higher vs non-AI · Shopify nearly 50% vs organic search", "vendor-reported, 4 · 5"],
            ["EU", "AI Act Art. 50 from 2026-08-02; ChatGPT VLOSE 2026-08-31, 159.1M", "filed, 2"],
            ["Enforcement", "none found on AI-answer ad disclosure, any jurisdiction", "filed"],
            ["Court record", "Microsoft 83-93% drops in click-through rates (ECF 1977-1)", "filed, 2"],
        ],
    },
    {
        "kind": "table",
        "title": "Supply side",
        "key": "A $1.8B valuation beside 2 of 34 vendors filing any revenue.",
        "cols": ["Metric", "Figure", "Label, tier"],
        "widths": [3.2, 6.6, 2.5],
        "rows": [
            ["Vendors rostered", "26 of 121 screened; 34 after Pass 3; 41 profiles", "analyst-derived"],
            ["Largest round", "Profound $180M Series D at $1.8B, 2026-09-15", "company-stated, 3"],
            ["Funding disclosed", "8 of 34", "company-stated, 5"],
            ["Price disclosed", "14 of 34; self-serve $20/mo to $999/month", "vendor-reported"],
            ["Filed revenue", "2 of 34; neither broken out to this product", "filed, 5 relayed"],
            ["Prompt set disclosed", "0 of 34 vendors; 0 of 41 profiles", "vendor-reported, 3"],
            ["Acquisitions", "Semrush → Adobe 2026-04-28; Scrunch → Sitecore 2026-06-03", "filed; company-stated"],
        ],
    },
    {
        "kind": "table",
        "title": "Demand: the 27-cell read",
        "key": "Willingness to pay: unknown in 27 of 27 cells. No buyer discloses a price paid.",
        "cols": ["Metric", "Figure", "Label, tier"],
        "widths": [3.2, 6.6, 2.5],
        "rows": [
            ["27-cell read", "7 spend · 1 attention · 11 none — checked · 8 blank", "analyst-derived, 5"],
            ["Signal carrying spend", "6 of 7 spend cells rest on S1 job postings", "company-stated, 3"],
            ["Where spend sits", "5 of 7 spend cells enterprise; 1 of 9 SMB cells", "analyst-derived, 3"],
            ["By sub-market", "organic 5 spend · paid 1 · agentic 1", "analyst-derived, 3"],
            ["Rule split", "strict none rule: 8 of 27 meet bar; loose: 27 of 27", "analyst-derived"],
        ],
    },
    {
        "kind": "bullets",
        "title": "Demand per vertical",
        "bullets": [
            "Skincare and beauty — 3 spend, all enterprise (e.l.f. AEO/GEO team, Direct Offers pilot, agentic role); 1 attention; 5 none",
            "Skincare agentic × enterprise: spend in customers/, attention in markets/ — both stand",
            "B2B SaaS — 3 spend, all organic (Actindo, AutoLeap, Pennylane and Mercury); 6 none",
            "High-CPA regulated — 1 spend (Cigna); 8 blank; no source states an SMB or mid-market band",
            "Floor read both ways: customers/ reads tier-3 postings as spend; source censuses read attention",
        ],
    },
    {
        "kind": "table",
        "title": "Proof: what the published cases show",
        "key": "No causal lift anywhere. H3 killed.",
        "cols": ["Metric", "Figure", "Label, tier"],
        "widths": [3.2, 6.6, 2.5],
        "rows": [
            ["Candidates screened", "~980, Passes 3 and 4", "measured-by-us"],
            ["Gold", "0; zero holdout, geo-split or switchback", "vendor-reported cases, 5"],
            ["Silver", "7 as graded in raw · 1 under grading rule 1", "vendor-reported cases, 5"],
            ["Direction of the Silvers", "2 of 7 negative or null (NerdWallet, TW3/Citead)", "filed 2; preprint 4"],
            ["Brand-side corroboration", "0 of 59 corroborate; 59 silent; ~107 unchecked", "measured-by-us, 3"],
            ["Paid-by-outcome", "0 yes / 0 no / 12 unknown", "vendor-reported, 6"],
            ["Anchor vertical", "skincare 0 Silver at ~133 screened", "vendor-reported, 5"],
        ],
    },
    {
        "kind": "table",
        "title": "The seven Silvers, graded two ways",
        "cols": ["Case", "Direction", "As graded", "Rule 1 (item missing)", "Tier"],
        "widths": [3.0, 4.0, 2.2, 2.2, 0.9],
        "rows": [
            ["Quattr / Men's Wearhouse", "up: 46% more clicks, treated pages", "Silver", "Bronze (3)", "5"],
            ["Sitefire / Pointhound", "up: +300% site visits from AI Search", "Silver", "Bronze (6)", "5"],
            ["Sitefire / Jerry", "up: +78% AI referral traffic", "Bronze / Silver", "Bronze (6)", "5"],
            ["Seer / SaaS HR client", "up: 300% increase in AI traffic", "Silver", "Bronze (3, 7)", "5"],
            ["Tiwari GSC audit", "mixed: 22% CTR drop", "Silver", "Bronze (3, 5)", "5"],
            ["NerdWallet 8-K", "negative: credit cards revenue decreased 24%", "Silver", "Bronze (6)", "2"],
            ["TW3 / Citead replication", "null: citation moved on none of ten families", "Silver", "Silver", "4"],
        ],
    },
    {
        "kind": "table",
        "title": "What brands that moved changed",
        "key": "H10, H11 confirmed at Pass 11; earlier marks stand beside (H10 not produced; H11 killed / not produced).",
        "cols": ["Kind of change", "Cases naming it", "Of which tier ≤3"],
        "widths": [4.5, 4.0, 3.8],
        "rows": [
            ["Content ops", "63", "16 (15 are llms.txt)"],
            ["Stack (vendor tool)", "58", "2"],
            ["Other (agency, PR, KPI, checkout)", "19", "1"],
            ["Org (hire, team, role)", "12", "12"],
            ["Paid media", "2 — both Google's Direct Offers pilot", "2"],
            ["None named", "9", "5"],
            ["Cases in set", "108; 99 name a change", "38"],
        ],
    },
    {
        "kind": "table",
        "title": "Hypotheses: 23, three readings, none averaged",
        "key": "Pass 9: 6 · 6 · 4 · 7. Register (review-1): 4 · 3 · 8 · 8. After Pass 11: 17 of 23 scored, 6 not produced.",
        "cols": ["ID", "Hypothesis", "Pass 9", "Register now"],
        "widths": [0.9, 5.4, 2.4, 3.6],
        "rows": [
            ["H3", "a vendor demonstrates causal lift", "killed", "killed"],
            ["H5", "corpus seeding measurably moves an answer", "killed", "unresolved — checked"],
            ["H6", "brand action precedes measured visibility change", "confirmed", "unresolved — checked"],
            ["H10", "org/content changes outnumber paid", "not produced", "Pass 11 confirmed"],
            ["H11", "tier-3 transition evidence per vertical", "killed", "not produced; Pass 11 confirmed"],
            ["H15", "composite-score vendors disclose prompt set", "killed", "unresolved — checked"],
            ["H16", "every published size is a forecast", "confirmed", "unresolved — checked"],
            ["HE2–HP4", "six panel hypotheses", "not produced", "not produced (Pass 10 held)"],
        ],
    },
    {
        "kind": "table",
        "title": "Programme done: 6 rows",
        "cols": ["Condition", "Bar", "Status"],
        "widths": [3.3, 3.2, 5.8],
        "rows": [
            ["Tier-3 share", "≥80%", "not met — 68.0% (17 of 25, Pass 9) · 44.4% (12 of 27, review-1)"],
            ["P1 engine × sub-market", "number or unknown — checked", "met — 9 of 9"],
            ["Segment matrix", "every cell read", "not met — 19 of 27; strict 8, loose 27"],
            ["Success stories", "Silver per vertical or absence", "met under both grade readings"],
            ["Hypotheses", "every H scored", "not met — 15 of 23 (register); 17 of 23 (after Pass 11)"],
            ["Pass 10", "3 predictions on panel", "not met — 0 of 3; held"],
        ],
    },
    {
        "kind": "bullets",
        "title": "What the evidence does not show",
        "bullets": [
            "Causal sales lift from any intervention (~980 screened)",
            "Whether AI referral converts above organic search (H1 unresolved)",
            "Any sub-market size, any price a buyer paid, any engine rate card",
            "A per-engine AI referral or conversion breakout from any retail publisher",
            "Brand-side confirmation of a vendor case: 0 of 59 checked; ~107 unchecked",
            "Neutral engine answers to buying-shaped prompts: six panel hypotheses not produced",
            "Anything behind blocked channels: sec.gov, Reddit, G2/Capterra, Indeed/Upwork, Google Trends",
        ],
    },
    {
        "kind": "bullets",
        "title": "Questions only the owner can answer",
        "note": "No lean is attached to any question.",
        "bullets": [
            "1. Do the seven plan-review-2 §8 appends (a–g) go into method/plan.md?",
            "2. Does Pass 10 reopen, and under which vantage — logged-in or out, which region or egress?",
            "3. Which Silver count governs: 7 as graded, or 1 under grading rule 1?",
            "4. Which sub-market does the hypothetical product target — organic, paid, agentic, or more than one?",
            "5. Does the builder-constraint question return to scope?",
        ],
    },
    {
        "kind": "bullets",
        "title": "Confidence and caveats",
        "bullets": [
            "Tier-3 share 68.0% and 44.4%, both below the 80% bar; shortfall is case corpus and forecasts",
            "Survivorship: published cases are winners; re-grading moved 12 up, 0 down",
            "Prices, counts, funding, cases: self-reports, none independently audited",
            "Measured by us: partial day-0 panel only, one date, one localised network path",
            "Conflicts kept side by side, never averaged: Silver 7 / 1; ad presence 26% / 4.47% / 0.00%",
            "Absence is only as strong as the channels checked; several were blocked, not exhausted",
            "Profiles go stale 2026-12-22; this deck carries evidence, not a verdict",
        ],
    },
]

# ---------------------------------------------------------------- helpers


def _text(frame, text, size, color=BODY, bold=False, align=PP_ALIGN.LEFT):
    frame.clear()
    p = frame.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Calibri"
    return p


def _box(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tb.text_frame.word_wrap = True
    return tb.text_frame


def title_bar(slide, title):
    bar = slide.shapes.add_shape(1, 0, 0, W, Inches(1.0))
    bar.fill.solid()
    bar.fill.fore_color.rgb = DARK
    bar.line.fill.background()
    tf = bar.text_frame
    tf.margin_left = Inches(0.5)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _text(tf, title, 28, WHITE, bold=True)
    accent = slide.shapes.add_shape(1, 0, Inches(1.0), W, Inches(0.06))
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT
    accent.line.fill.background()


def footer(slide, n):
    tf = _box(slide, Inches(0.5), Inches(7.0), Inches(10.5), Inches(0.35))
    _text(tf, FOOTER, 10, GREY)
    tf2 = _box(slide, Inches(11.3), Inches(7.0), Inches(1.6), Inches(0.35))
    _text(tf2, str(n), 10, GREY, align=PP_ALIGN.RIGHT)


def key_line(slide, text, y):
    tf = _box(slide, Inches(0.5), y, Inches(12.3), Inches(0.5))
    _text(tf, text, 15, ACCENT, bold=True)


def bullets(slide, items, y, size=18):
    tf = _box(slide, Inches(0.7), y, Inches(12.0), Inches(5.5))
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = ("• " if not item[:2].rstrip(".").isdigit() else "") + item
        r.font.size = Pt(size)
        r.font.color.rgb = BODY
        r.font.name = "Calibri"
        p.space_after = Pt(10)


def table(slide, cols, rows, widths, y):
    assert len(rows) <= 8, "keep tables to 8 rows or fewer"
    n_rows = len(rows) + 1
    height = Inches(0.42) * n_rows
    shape = slide.shapes.add_table(n_rows, len(cols), Inches(0.5), y, Inches(sum(widths)), height)
    tbl = shape.table
    for j, w in enumerate(widths):
        tbl.columns[j].width = Inches(w)
    for j, c in enumerate(cols):
        cell = tbl.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT
        _text(cell.text_frame, c, 13, WHITE, bold=True)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = ROW_ALT if i % 2 == 0 else WHITE
            _text(cell.text_frame, val, 11, BODY, bold=(j == 0))


# ---------------------------------------------------------------- build


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    for n, s in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(blank)
        if s["kind"] == "title":
            bg = slide.shapes.add_shape(1, 0, 0, W, Inches(3.6))
            bg.fill.solid()
            bg.fill.fore_color.rgb = DARK
            bg.line.fill.background()
            tf = _box(slide, Inches(0.7), Inches(1.2), Inches(12), Inches(1.4))
            _text(tf, s["title"], 36, WHITE, bold=True)
            tf = _box(slide, Inches(0.7), Inches(2.6), Inches(12), Inches(0.6))
            _text(tf, s["subtitle"], 18, WHITE)
            bar = slide.shapes.add_shape(1, 0, Inches(3.6), W, Inches(0.08))
            bar.fill.solid()
            bar.fill.fore_color.rgb = ACCENT
            bar.line.fill.background()
            bullets(slide, s["lines"], Inches(4.2), size=18)
            tf = _box(slide, Inches(0.7), Inches(6.3), Inches(12), Inches(0.4))
            _text(tf, SOURCE, 11, GREY)
            footer(slide, n)
            continue
        title_bar(slide, s["title"])
        y = Inches(1.3)
        if s.get("quote"):
            tf = _box(slide, Inches(0.7), y, Inches(12), Inches(1.2))
            _text(tf, s["quote"], 17, DARK, bold=True)
            y = Inches(2.6)
        if s.get("key"):
            key_line(slide, s["key"], y)
            y = y + Inches(0.6)
        if s.get("note"):
            key_line(slide, s["note"], y)
            y = y + Inches(0.6)
        if s["kind"] == "table":
            table(slide, s["cols"], s["rows"], s["widths"], y)
        else:
            bullets(slide, s["bullets"], y)
        footer(slide, n)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
